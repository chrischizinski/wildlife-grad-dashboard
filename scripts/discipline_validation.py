#!/usr/bin/env python3
"""
Blind human validation of discipline labels.

The pipeline's "gold" labels are auto-seeded from its own rule output, so they
cannot measure accuracy. This tool builds an independent check:

  sample: draw postings and write a CSV for a human to label WITHOUT seeing the
          pipeline's label (seeing it anchors the answer). Small classes are
          labeled in full; the rest is a seeded random sample. The pipeline's
          labels go to a separate key file that the reviewer does not need.
  score:  compare the filled-in CSV with the key and report stratified accuracy
          with a 95% confidence interval, per-label precision, and confusions.

Reviewer instructions: open the .xlsx and pick an answer from the dropdown in
`human_discipline` (a CSV also works: type the label, case-insensitive, unique
prefix is fine, e.g. "fish" or "wild"). Use `not_graduate` if the posting is not
a graduate position, `unclear` if you cannot tell from the posting.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import random
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

DISCIPLINES = [
    "Environmental Sciences",
    "Fisheries and Aquatic",
    "Wildlife",
    "Entomology",
    "Forestry and Habitat",
    "Agriculture",
    "Human Dimensions",
    "Other",
]
NOT_GRAD = "not_graduate"
UNCLEAR = "unclear"
ANSWERS = DISCIPLINES + [NOT_GRAD, UNCLEAR]

DEFAULT_POSITIONS = Path("web/data/dashboard_positions.json")
DEFAULT_SAMPLE = Path("data/validation/discipline_sample.xlsx")
DEFAULT_KEY = Path("data/validation/discipline_sample_key.json")
DEFAULT_CORRECTIONS = Path("data/validation/label_corrections.json")
EXCERPT_CHARS = 700
SAMPLE_COLUMNS = [
    "position_key",
    "title",
    "organization",
    "location",
    "salary",
    "url",
    "description_excerpt",
    "human_discipline",
    "notes",
]


def position_key(row: Dict[str, Any]) -> str:
    return str(row.get("url") or f"{row.get('title')}|{row.get('organization')}")


def normalize_answer(value: Any) -> Optional[str]:
    """Map reviewer text to an allowed answer; None if blank or ambiguous."""
    text = re.sub(r"\s+", " ", str(value or "")).strip().lower()
    if not text:
        return None
    exact = [a for a in ANSWERS if a.lower() == text]
    if exact:
        return exact[0]
    prefix = [a for a in ANSWERS if a.lower().startswith(text)]
    return prefix[0] if len(prefix) == 1 else None


def draw_sample(
    rows: List[Dict[str, Any]], n: int, rare_below: int, seed: int
) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Census of rare predicted classes plus a random sample of the rest."""
    counts = Counter(r["discipline_primary"] for r in rows)
    rare_labels = {label for label, c in counts.items() if c < rare_below}
    rare = [r for r in rows if r["discipline_primary"] in rare_labels]
    common = [r for r in rows if r["discipline_primary"] not in rare_labels]
    n_common = max(0, min(len(common), n - len(rare)))
    rng = random.Random(seed)
    picked = rng.sample(common, n_common)
    meta = {
        "seed": seed,
        "rare_labels": sorted(rare_labels),
        "strata": {
            "rare": {"population": len(rare), "sampled": len(rare)},
            "common": {"population": len(common), "sampled": n_common},
        },
    }
    sample = [(r, "rare") for r in rare] + [(r, "common") for r in picked]
    rng.shuffle(sample)
    return sample, meta


def excerpt(row: Dict[str, Any]) -> str:
    text = re.sub(r"\s+", " ", str(row.get("description") or "")).strip()
    return text[:EXCERPT_CHARS] + ("..." if len(text) > EXCERPT_CHARS else "")


def wilson(correct: int, total: int, z: float = 1.96) -> Tuple[float, float]:
    if total == 0:
        return 0.0, 1.0
    p = correct / total
    denom = 1 + z * z / total
    centre = (p + z * z / (2 * total)) / denom
    half = z * math.sqrt(p * (1 - p) / total + z * z / (4 * total * total)) / denom
    return max(0.0, centre - half), min(1.0, centre + half)


def stratified_accuracy(
    outcomes: Dict[str, List[bool]], population: Dict[str, int]
) -> Tuple[float, float]:
    """Accuracy and 95% half-width from per-stratum correct/incorrect lists."""
    total = sum(population.values())
    estimate = 0.0
    variance = 0.0
    for stratum, results in outcomes.items():
        m = len(results)
        if m == 0:
            continue
        weight = population[stratum] / total
        p = sum(results) / m
        estimate += weight * p
        if m > 1:
            fpc = 1 - m / population[stratum]
            variance += weight**2 * fpc * p * (1 - p) / (m - 1)
    return estimate, 1.96 * math.sqrt(variance)


def score(key: Dict[str, Any], answers: Dict[str, str]) -> Dict[str, Any]:
    outcomes: Dict[str, List[bool]] = defaultdict(list)
    per_label: Dict[str, List[bool]] = defaultdict(list)
    confusions: Counter = Counter()
    not_grad = Counter()
    unclear = 0
    unanswered = 0
    for pos_key, info in key["rows"].items():
        answer = answers.get(pos_key)
        if answer is None:
            unanswered += 1
            continue
        if answer == UNCLEAR:
            unclear += 1
            continue
        stratum = info["stratum"]
        if answer == NOT_GRAD:
            not_grad[stratum] += 1
            continue
        correct = answer == info["predicted"]
        outcomes[stratum].append(correct)
        per_label[info["predicted"]].append(correct)
        if not correct:
            confusions[(info["predicted"], answer)] += 1

    population = {s: v["population"] for s, v in key["meta"]["strata"].items()}
    estimate, half = stratified_accuracy(outcomes, population)
    labeled = sum(len(v) for v in outcomes.values())
    return {
        "labeled": labeled,
        "unanswered": unanswered,
        "unclear": unclear,
        "not_graduate": dict(not_grad),
        "accuracy": estimate,
        "ci95_half_width": half,
        "per_label": {
            label: {
                "n": len(res),
                "correct": sum(res),
                "wilson95": wilson(sum(res), len(res)),
            }
            for label, res in sorted(per_label.items())
        },
        "top_confusions": confusions.most_common(10),
    }


def write_workbook(records: List[Dict[str, Any]], path: Path) -> None:
    """Write the labeling workbook: dropdown answers, wrapped text, progress count."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Label"
    lists = wb.create_sheet("Lists")
    for i, answer in enumerate(ANSWERS, start=1):
        lists.cell(row=i, column=1, value=answer)
    lists.sheet_state = "hidden"

    ws.append(SAMPLE_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for record in records:
        ws.append([record.get(col, "") for col in SAMPLE_COLUMNS])

    widths = {"position_key": 8, "title": 45, "organization": 28, "location": 28,
              "salary": 16, "url": 26, "description_excerpt": 80,
              "human_discipline": 26, "notes": 36}
    for idx, col in enumerate(SAMPLE_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths[col]
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap

    label_col = get_column_letter(SAMPLE_COLUMNS.index("human_discipline") + 1)
    last_row = len(records) + 1
    validation = DataValidation(
        type="list",
        formula1=f"=Lists!$A$1:$A${len(ANSWERS)}",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Pick from the list",
        error="Choose one of the dropdown answers.",
    )
    ws.add_data_validation(validation)
    validation.add(f"{label_col}2:{label_col}{last_row}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SAMPLE_COLUMNS))}{last_row}"

    info = wb.create_sheet("Instructions", 0)
    lines = [
        "Discipline label validation",
        "",
        "1. Go to the 'Label' sheet.",
        "2. For each posting pick an answer from the dropdown in 'human_discipline'.",
        "3. Label from the posting itself. The pipeline's label is deliberately not shown.",
        "4. 'not_graduate' = not a graduate position; 'unclear' = cannot tell from the posting.",
        "5. Use 'notes' when a posting could fit two disciplines.",
        "6. Save the file when done, then run: python scripts/discipline_validation.py score",
        "",
        f'=\"Labeled so far: \"&COUNTA(Label!{label_col}2:{label_col}{last_row})&\" of {len(records)}\"',
    ]
    for i, line in enumerate(lines, start=1):
        info.cell(row=i, column=1, value=line)
    info["A1"].font = Font(bold=True, size=14)
    info.column_dimensions["A"].width = 90
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def read_answers(path: Path) -> List[Dict[str, Any]]:
    """Read labeling rows from the workbook (sheet 'Label') or a CSV."""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        sheet = load_workbook(path, data_only=True)["Label"]
        rows = list(sheet.iter_rows(values_only=True))
        header = [str(h) for h in rows[0]]
        return [
            {h: ("" if v is None else v) for h, v in zip(header, r)}
            for r in rows[1:]
            if any(v is not None for v in r)
        ]
    with open(path, encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def apply_corrections(answers: Dict[str, str], path: Path) -> Tuple[int, int]:
    """Apply reviewer corrections in place; a null label excludes the row."""
    if not path.exists():
        return 0, 0
    corrections = json.loads(path.read_text(encoding="utf-8"))["corrections"]
    applied = excluded = 0
    for pos_key, entry in corrections.items():
        if entry.get("label") is None:
            if answers.pop(pos_key, None) is not None:
                excluded += 1
        elif pos_key in answers:
            answers[pos_key] = entry["label"]
            applied += 1
    return applied, excluded


def cmd_sample(args: argparse.Namespace) -> int:
    rows = json.loads(args.positions.read_text(encoding="utf-8"))
    rows = rows["positions"] if isinstance(rows, dict) else rows
    sample, meta = draw_sample(rows, args.n, args.rare_below, args.seed)

    records = [
        {
            "position_key": position_key(row),
            "title": row.get("title", ""),
            "organization": row.get("organization", ""),
            "location": row.get("location", ""),
            "salary": row.get("salary", ""),
            "url": row.get("url", ""),
            "description_excerpt": excerpt(row),
            "human_discipline": "",
            "notes": "",
        }
        for row, _stratum in sample
    ]
    args.sample.parent.mkdir(parents=True, exist_ok=True)
    if args.sample.suffix.lower() == ".xlsx":
        write_workbook(records, args.sample)
    else:
        with open(args.sample, "w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=SAMPLE_COLUMNS)
            writer.writeheader()
            writer.writerows(records)
    key = {
        "meta": meta,
        "rows": {
            position_key(r): {
                "predicted": r["discipline_primary"],
                "source": r.get("discipline_refinement_source", ""),
                "stratum": stratum,
            }
            for r, stratum in sample
        },
    }
    args.key.write_text(json.dumps(key, indent=2), encoding="utf-8")
    print(f"Wrote {len(sample)} rows to {args.sample}")
    print(f"Key (do not open before labeling) -> {args.key}")
    print(f"Strata: {meta['strata']}")
    print("Allowed answers: " + ", ".join(ANSWERS))
    return 0


def cmd_score(args: argparse.Namespace) -> int:
    key = json.loads(args.key.read_text(encoding="utf-8"))
    answers: Dict[str, str] = {}
    bad = []
    for row in read_answers(args.sample):
        raw = row.get("human_discipline", "")
        if not str(raw).strip():
            continue
        normalized = normalize_answer(raw)
        if normalized is None:
            bad.append((row["position_key"], raw))
        else:
            answers[row["position_key"]] = normalized
    if bad:
        print(f"{len(bad)} unrecognized answers (fix these rows):", file=sys.stderr)
        for pos_key, raw in bad[:10]:
            print(f"  {raw!r}  {pos_key}", file=sys.stderr)
        return 1

    applied, excluded = apply_corrections(answers, args.corrections)
    if applied or excluded:
        print(f"Corrections applied: {applied}; excluded pending: {excluded}")
    result = score(key, answers)
    print(f"Labeled: {result['labeled']}  unclear: {result['unclear']}  "
          f"not_graduate: {result['not_graduate']}  unanswered: {result['unanswered']}")
    print(
        f"Stratified accuracy: {result['accuracy']:.1%} "
        f"± {result['ci95_half_width']:.1%} (95% CI)"
    )
    print("Per predicted label (n, correct, Wilson 95%):")
    for label, info in result["per_label"].items():
        lo, hi = info["wilson95"]
        print(f"  {label:24} n={info['n']:3} correct={info['correct']:3} "
              f"[{lo:.0%}, {hi:.0%}]")
    print("Most common errors (pipeline -> human):")
    for (pred, human), count in result["top_confusions"]:
        print(f"  {count:3}  {pred} -> {human}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    sub = parser.add_subparsers(dest="command", required=True)

    p_sample = sub.add_parser("sample", help="write the blind labeling CSV")
    p_sample.add_argument("--n", type=int, default=120)
    p_sample.add_argument("--rare-below", type=int, default=15,
                          help="label every row of classes with fewer rows than this")
    p_sample.add_argument("--seed", type=int, default=20261002)
    p_sample.add_argument("--positions", type=Path, default=DEFAULT_POSITIONS)
    p_sample.add_argument("--sample", type=Path, default=DEFAULT_SAMPLE)
    p_sample.add_argument("--key", type=Path, default=DEFAULT_KEY)
    p_sample.set_defaults(func=cmd_sample)

    p_score = sub.add_parser("score", help="score a filled-in CSV")
    p_score.add_argument("--sample", type=Path, default=DEFAULT_SAMPLE)
    p_score.add_argument("--key", type=Path, default=DEFAULT_KEY)
    p_score.add_argument("--corrections", type=Path, default=DEFAULT_CORRECTIONS)
    p_score.set_defaults(func=cmd_score)

    args = parser.parse_args()
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
